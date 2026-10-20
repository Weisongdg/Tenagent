"""使用 openpyxl 读取当前 12 列问题模板；仅做字段映射和来源记录。"""

from contextlib import closing
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from config import PROJECT_DIR

from .schemas import ExcelReadResult, RawIssue

DEFAULT_FILE = "Process_Issue_Tracking.xlsx"
DEFAULT_SHEET = "ProcessIssueTracking"

# 键为 Excel 原表头；仅在匹配表头时忽略首尾空白和大小写。
HEADER_FIELDS = {
    "No": "no",
    "Prod": "product",
    "When": "when",
    "Area": "area",
    "station": "station",
    "problem description": "problem_description",
    "category": "category",
    "Impact Qty": "impact_qty",
    "Root cause & Possible reason": "root_cause",
    "Owner": "owner",
    "Status": "status",
    "Action": "action",
}


class ExcelReadError(ValueError):
    """文件、工作表或模板不符合要求时，给调用方明确的失败原因。"""


def _map_headers(headers: tuple) -> dict[str, int]:
    expected = {header.strip().casefold(): field for header, field in HEADER_FIELDS.items()}
    columns: dict[str, int] = {}
    for index, value in enumerate(headers):
        if value is None or value == "":
            continue
        key = str(value).strip().casefold()
        if key not in expected:
            raise ExcelReadError(f"不支持的表头：{value}。当前读取器支持约定的 12 列问题模板。")
        field = expected[key]
        if field in columns:
            raise ExcelReadError(f"存在重复表头：{value}。请保证每个字段只出现一次。")
        columns[field] = index

    missing = [header for header, field in HEADER_FIELDS.items() if field not in columns]
    if missing:
        raise ExcelReadError(f"第 1 行缺少必需表头：{', '.join(missing)}。")
    return columns


def read_issues(file_path: str | Path, sheet_name: str = DEFAULT_SHEET) -> ExcelReadResult:
    """读取指定工作表；相对路径基于项目根目录，返回全部非空行。"""
    path = Path(file_path).expanduser()
    if not path.is_absolute():
        path = PROJECT_DIR / path
    path = path.resolve()
    if not path.is_file():
        raise ExcelReadError(f"Excel 文件不存在或不是普通文件：{path}")
    if path.suffix.lower() != ".xlsx":
        raise ExcelReadError("当前只支持 .xlsx 文件。")

    try:
        # closing 在成功或失败时都会关闭工作簿；本功能从不调用 save。
        with closing(load_workbook(path, read_only=True, data_only=False)) as workbook:
            if sheet_name not in workbook.sheetnames:
                raise ExcelReadError(
                    f"工作表不存在：{sheet_name}。可用工作表：{', '.join(workbook.sheetnames)}"
                )
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            columns = _map_headers(next(rows, ()))
            mapped_columns = set(columns.values())
            issues = []
            for row_number, row in enumerate(rows, start=2):
                if all(value is None or value == "" for value in row):
                    continue
                if any(
                    index not in mapped_columns and value is not None and value != ""
                    for index, value in enumerate(row)
                ):
                    raise ExcelReadError(f"第 {row_number} 行存在没有表头的数据列，请先补充表头。")
                values = {field: row[index] for field, index in columns.items()}
                issues.append(RawIssue(source_row=row_number, **values))

            return ExcelReadResult(
                source_file=str(path),
                sheet_name=sheet_name,
                column_mapping={
                    field: get_column_letter(index + 1) for field, index in columns.items()
                },
                issues=issues,
            )
    except ExcelReadError:
        raise
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError) as error:
        raise ExcelReadError(
            f"无法读取 Excel 文件：{path.name}（{type(error).__name__}）。"
            "请检查文件是否有效、是否加密以及是否具有读取权限。"
        ) from error
